"""
Document conversion: HTML/Markdown/TXT → PDF, CSV → Excel.
Dependencies: weasyprint, markdown, openpyxl.
"""
import csv
import io


def html_to_pdf(data: bytes, meta: dict) -> tuple[bytes, str]:
    from weasyprint import HTML
    html_str = data.decode("utf-8", errors="replace")
    pdf_bytes = HTML(string=html_str).write_pdf()
    return pdf_bytes, "output.pdf"


def markdown_to_pdf(data: bytes, meta: dict) -> tuple[bytes, str]:
    import markdown as md_lib
    from weasyprint import HTML

    md_text = data.decode("utf-8", errors="replace")
    body_html = md_lib.markdown(md_text, extensions=["tables", "fenced_code", "nl2br"])
    styled = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: Georgia, serif; max-width: 780px; margin: 48px auto; padding: 0 24px;
         line-height: 1.65; color: #222; font-size: 15px; }}
  h1, h2, h3, h4 {{ color: #111; margin-top: 1.4em; }}
  pre {{ background: #f5f5f5; padding: 12px; border-radius: 4px; overflow-x: auto;
        font-size: 13px; line-height: 1.4; }}
  code {{ background: #f5f5f5; padding: 2px 5px; border-radius: 3px; font-size: 13px; }}
  table {{ border-collapse: collapse; width: 100%; margin: 1em 0; }}
  th, td {{ border: 1px solid #ccc; padding: 8px 12px; text-align: left; }}
  th {{ background: #f0f0f0; font-weight: 600; }}
  blockquote {{ border-left: 3px solid #ccc; margin: 0; padding-left: 16px; color: #555; }}
</style>
</head>
<body>{body_html}</body>
</html>"""
    pdf_bytes = HTML(string=styled).write_pdf()
    return pdf_bytes, "output.pdf"


def txt_to_pdf(data: bytes, meta: dict) -> tuple[bytes, str]:
    from weasyprint import HTML

    text = data.decode("utf-8", errors="replace")
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: "Courier New", monospace; white-space: pre-wrap;
         padding: 48px; line-height: 1.55; font-size: 13px; color: #222; }}
</style>
</head>
<body>{escaped}</body>
</html>"""
    pdf_bytes = HTML(string=html).write_pdf()
    return pdf_bytes, "output.pdf"


def csv_to_excel(data: bytes, meta: dict) -> tuple[bytes, str]:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)

    wb = openpyxl.Workbook()
    ws = wb.active

    header_fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")

    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 1:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=False)

    # Auto-fit column widths (approximate)
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "output.xlsx"


# ─── Dispatcher ───────────────────────────────────────────────────────────────

DOCUMENT_PROCESSORS = {
    "html-to-pdf":     html_to_pdf,
    "markdown-to-pdf": markdown_to_pdf,
    "txt-to-pdf":      txt_to_pdf,
    "csv-to-excel":    csv_to_excel,
}
